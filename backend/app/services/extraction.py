"""Text and structure extraction for uploaded omics documents.

The entry point is :func:`extract_text`, which routes a stored file to the
right extractor based on its extension and returns an :class:`ExtractionResult`.
Structured omics files (VCF/BED/tabular matrices) are also sniffed by
``app.services.omics_parser`` so the caller can persist machine-readable
content into ``Document.raw_extracted_json``.
"""

import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from email import message_from_bytes
from email.policy import default
from pathlib import Path

import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader

from app.services.omics_parser import parse_omics, read_text


class UnsupportedFormat(Exception):
    """The file's inner format cannot be parsed yet (e.g. .xlsx)."""


@dataclass
class ExtractionResult:
    kind: str  # pdf | image | docx | email | vcf | bed | matrix | text
    text: str = ""
    structure: dict | None = None


def extract_pdf(file: str) -> str:
    reader = PdfReader(file)
    pages = reader.pages
    texts = "\n".join(page.extract_text() for page in pages)

    return texts


def extract_img(file: str) -> str:
    image = Image.open(file)
    text = pytesseract.image_to_string(image)

    return text


def extract_docx(file: str) -> str:
    docs = Document(file)
    text = "\n".join(doc.text for doc in docs.paragraphs)
    table = ""
    for tables in docs.tables:
        for row in tables.rows:
            table += f"{[cell.text for cell in row.cells]}\n"

    return text + f"\n{table}"


def extract_email(file: str) -> str:
    with open(file, "rb") as f:
        message = message_from_bytes(f.read(), policy=default)

    body = message.get_body(preferencelist=("plain"))
    text = body.get_content() if body else ""

    return text


DOCUMENT_EXTRACTORS: dict[str, Callable[[str], str]] = {
    "pdf": extract_pdf,
    "image": extract_img,
    "docx": extract_docx,
    "email": extract_email,
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
TEXT_EXTENSIONS = {".txt", ".tsv", ".csv", ".vcf", ".bed"}


def _sniff_zip(file: str) -> str:
    with zipfile.ZipFile(file) as archive:
        names = archive.namelist()
        if any(name.startswith("word/") for name in names):
            return "docx"
        if any(name.startswith("xl/") for name in names):
            return "xlsx"
    return "unknown"


def extract_text(file: str) -> ExtractionResult:
    """Route a stored file to the right extractor and return its content."""
    path = Path(file)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return ExtractionResult("pdf", extract_pdf(file))
    if ext in IMAGE_EXTENSIONS:
        return ExtractionResult("image", extract_img(file))
    if ext == ".docx":
        return ExtractionResult("docx", extract_docx(file))
    if ext == ".eml":
        return ExtractionResult("email", extract_email(file))
    if ext == ".zip":
        inner = _sniff_zip(file)
        if inner == "docx":
            return ExtractionResult("docx", extract_docx(file))
        raise UnsupportedFormat(
            f"ZIP contents look like '{inner}'. Only .docx inside .zip is supported."
        )

    text = read_text(path)
    structure = parse_omics(text)
    if structure["kind"] in {"vcf", "bed", "matrix"}:
        return ExtractionResult(structure["kind"], text, structure)
    return ExtractionResult("text", text, None)


def main() -> None:
    """CLI smoke test: python -m app.services.extraction <file>"""
    import json
    import sys

    result = extract_text(sys.argv[1])
    print(json.dumps({"kind": result.kind, "structure": result.structure}, indent=2))
    if result.text:
        print("\n--- text preview ---")
        print("\n".join(result.text.splitlines()[:10]))


if __name__ == "__main__":
    main()